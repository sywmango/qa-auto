# 단위테스트 결과를 회사 UT 양식 엑셀(Dev 조치사항 시트)에 옮겨 적습니다. (메뉴 공통: 테스트 제목이 'AD230_TC_001 …' 처럼 화면ID_TC/EC_번호로 시작하면 됨)
# 사용: python scripts/write-results.py <원본.xlsx> <새파일.xlsx> <결과json 폴더> [수동결과.json]
#  - 결과json 폴더: playwright 의 --reporter=json 결과 파일들 (나중 파일이 앞 파일을 덮어씀)
#  - 수동결과.json: { "AD230_EC_001": {"result": "미수행", "text": "사유", "data": ""}, ... }
# 적는 칸: H(QA) I(결과 정상/오류) J(내용) K(날짜) L(첨부파일 = 증적 캡처 파일명)
# 회사 양식(v1)에 없는 열은 만들지 않음 (v2~v11 에 있던 Q '테스트 데이터 / 증적' 열은 지움)
import json, os, re, sys, glob, datetime
import openpyxl
from openpyxl.styles import Alignment

src, dst, res_dir = sys.argv[1], sys.argv[2], sys.argv[3]
manual = json.load(open(sys.argv[4], encoding='utf8')) if len(sys.argv) > 4 else {}

if os.path.exists(dst):
    try:
        os.rename(dst, dst)  # 엑셀에서 열려 있으면 여기서 실패
    except OSError:
        sys.exit(f'{dst} 가 열려 있습니다. 엑셀을 닫고 다시 실행하세요.')

ID_RE = r'([A-Z]{2}\d{3}_(?:TC|EC)_\d+)'
is_id = lambda v: bool(v) and re.fullmatch(ID_RE, str(v)) is not None
results = {}
merged = {}


def walk(suite):
    for s in suite.get('suites', []):
        yield from walk(s)
    for spec in suite.get('specs', []):
        yield spec


for f in sorted(glob.glob(os.path.join(res_dir, '*.json')), key=os.path.getmtime):
    rep = json.load(open(f, encoding='utf8'))
    for top in rep.get('suites', []):
        for spec in walk(top):
            m = re.match(ID_RE, spec['title'])
            if not m:
                continue
            for t in spec.get('tests', []):
                if not t.get('results'):
                    continue
                r = t['results'][-1]
                anns = t.get('annotations', []) + r.get('annotations', [])
                notes = list(dict.fromkeys(a.get('description', '') for a in anns if a.get('type') == '실제결과'))
                skip = list(dict.fromkeys(a.get('description', '') for a in anns if a.get('type') in ('미수행', 'skip')))
                status = r['status']
                err = (r.get('error') or {}).get('message', '')
                err = re.sub(r'\x1b\[[0-9;]*m', '', err).split('\n')[0][:200]
                when = r.get('startTime', '')[:10]
                entry = {
                    'result': '정상' if status == 'passed' else '미수행' if status == 'skipped' else '오류',
                    'text': '\n'.join(notes + skip + ([f'오류: {err}'] if status not in ('passed', 'skipped') and err else [])),
                    'date': when,
                    'attach': '',  # 실패 스크린샷은 test-results 가 실행마다 지워지므로 증적은 수동결과.json 의 attach 로 지정
                }
                key = m.group(1)
                # 한 ID 를 여러 테스트로 나눈 경우(같은 결과 파일 안) 합침: 하나라도 오류면 오류
                prev = merged.get(key)
                if prev and prev.get('_file') == f:
                    rank = {'오류': 2, '미수행': 1, '정상': 0}
                    entry['result'] = max(prev['result'], entry['result'], key=lambda x: rank.get(x, 0))
                    entry['text'] = prev['text'] + '\n' + entry['text']
                entry['_file'] = f
                merged[key] = entry
                results[key] = entry

results.update({k: {**results.get(k, {}), **v} for k, v in manual.items()})

wb = openpyxl.load_workbook(src)
ws = wb['Dev 조치사항']
today = datetime.date.today().isoformat()
done = 0
for row in range(6, ws.max_row + 1):
    tid = ws.cell(row, 4).value
    if not is_id(tid) or tid not in results:
        continue
    r = results[tid]
    ws.cell(row, 8).value = '최신영'
    # 엑셀 결과 칸은 정상/오류만: 개선은 정상으로 두고 내용 앞에 [개선], 미수행은 결과를 비우고 내용 앞에 [미수행]
    res = r.get('result', '')
    excel_res = {'개선': '정상', '확인필요': '정상', '미수행': '', '미실행': ''}.get(res, res)
    prefix = {'개선': '[개선] ', '확인필요': '[개선] ', '미수행': '[미수행] '}.get(res, '')
    ws.cell(row, 9).value = excel_res
    # 정상은 QA내용을 비움 (개선으로 정상 처리한 행은 [개선] 내용을 남김)
    ws.cell(row, 10).value = '' if res == '정상' else prefix + r.get('text', '')
    # 날짜는 글자가 아니라 날짜 값으로 (테스트 케이스 시트의 MAXIFS·INDEX 수식이 이 칸을 날짜로 비교함)
    ws.cell(row, 11).value = datetime.datetime.strptime(r.get('date') or today, '%Y-%m-%d')
    ws.cell(row, 11).number_format = 'yyyy-mm-dd'
    ws.cell(row, 12).value = r.get('attach', '')
    for c in (10,):
        ws.cell(row, c).alignment = Alignment(wrap_text=True, vertical='top')
    done += 1

# 양식에 없던 Q 열(테스트 데이터 / 증적)이 원본에 있으면 지움
if ws.max_column >= 17 and str(ws.cell(5, 17).value or '').startswith('테스트 데이터'):
    ws.delete_cols(17)
    ws.column_dimensions.pop('Q', None)
# 수식 결과(캐시)가 비어 있으므로 엑셀이 열 때 다시 계산하게 함
wb.calculation.fullCalcOnLoad = True
wb.save(dst)

# 보고서용 요약 JSON (결과 + 테스트 케이스 시트의 케이스/기대결과)
tc = wb['테스트 케이스']
cases = []
for row in range(6, tc.max_row + 1):
    tid = tc.cell(row, 4).value
    if not is_id(tid):
        continue
    cases.append({'id': tid, 'screen': tc.cell(row, 2).value or '', 'func': tc.cell(row, 3).value or '',
                  'case': tc.cell(row, 5).value or '', 'step': tc.cell(row, 6).value or '', 'expected': tc.cell(row, 7).value or '',
                  **results.get(tid, {'result': '미실행', 'text': ''})})
with open(os.path.splitext(dst)[0] + '_결과.json', 'w', encoding='utf8') as f:
    json.dump(cases, f, ensure_ascii=False, indent=1)
missing = [ws.cell(r, 4).value for r in range(6, ws.max_row + 1) if is_id(ws.cell(r, 4).value) and ws.cell(r, 4).value not in results]
print(f'{done}건 기록 → {dst}')
print('결과 없는 케이스:', ', '.join(missing) or '없음')
