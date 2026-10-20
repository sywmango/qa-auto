# 공통 UT 양식(테스트 케이스 + Dev 조치사항 시트)에 메뉴의 테스트 케이스를 채워 새 엑셀을 만듭니다. (메뉴 공통)
# 사용: python scripts/build-ut-sheet.py <공통양식.xlsx> <cases.json> <새파일.xlsx>
#  - cases.json: {"menu": "정기점검계획_AD210", "common": [[탭, 기능, ID, CASE, STEP, 기대결과], ...], "edge": [...]}
#    탭이 "" 이면 바로 위 행과 같은 탭 (양식처럼 셀 병합)
#  - 양식의 열·시트는 그대로 두고, 6행부터 내용만 바꿈. H~K 수식은 양식 6행 수식을 행 번호만 바꿔 넣음
import copy, json, re, sys
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

src, cases_path, dst = sys.argv[1:4]
cases = json.load(open(cases_path, encoding='utf8'))
wb = openpyxl.load_workbook(src)
tc, dev = wb['테스트 케이스'], wb['Dev 조치사항']


def find_row(ws, text):
    return next(r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == text)


def snapshot(ws, row):
    return [(copy.copy(ws.cell(row, c)._style), ws.cell(row, c).value) for c in range(1, ws.max_column + 1)], ws.row_dimensions[row].height


def shift(v, r):
    """양식 6행 수식의 D6 / $D6 을 r 행으로"""
    if isinstance(v, ArrayFormula):
        return ArrayFormula(re.sub(r'\d+$', str(r), v.ref), shift(v.text, r))
    if isinstance(v, str) and v.startswith('='):
        return re.sub(r'(?<![$\w])(\$?)D6(?!\d)', lambda m: f'{m.group(1)}D{r}', v)
    return v


plans = {}
for ws in (tc, dev):
    edge = find_row(ws, '엣지케이스')
    plans[ws.title] = {'data': snapshot(ws, 6), 'title': snapshot(ws, edge), 'head': snapshot(ws, edge + 1),
                       'blank': snapshot(ws, edge - 1), 'width': ws.max_column}
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= 6:
            ws.unmerge_cells(str(m))
    ws.delete_rows(6, ws.max_row)


def put(ws, r, snap, values=None, formulas=False):
    cells, h = snap
    for c, (style, v) in enumerate(cells, 1):
        cell = ws.cell(r, c)
        cell._style = copy.copy(style)
        if values is not None and c <= len(values):
            cell.value = values[c - 1]
        elif formulas and c > 4:
            cell.value = shift(v, r)
        elif values is None:
            cell.value = v
        else:
            cell.value = None
    ws.row_dimensions[r].height = h


def merge_runs(ws, col, start, end, key):
    r = start
    while r <= end:
        e = r
        while e + 1 <= end and key(e + 1) == key(r):
            e += 1
        if e > r:
            ws.merge_cells(start_row=r, start_column=col, end_row=e, end_column=col)
        r = e + 1


def write(ws, is_tc):
    p = plans[ws.title]
    menu = cases['menu']
    r = 6
    blocks = []
    for name in ('common', 'edge'):
        rows = cases[name]
        if name == 'edge':
            r += 2  # 양식처럼 두 줄 띄움
            put(ws, r, p['title'])
            put(ws, r + 1, p['head'])
            r += 2
        start = r
        tab = ''
        tabs, funcs = [], []
        for t, func, tid, case, step, exp in rows:
            tab = t or tab
            vals = [menu, tab, func, tid] + ([case, step, exp] if is_tc else [])
            put(ws, r, p['data'], formulas=True)
            for c, v in enumerate(vals, 1):
                ws.cell(r, c).value = v
            # 병합되는 아래 칸은 비움 (값은 첫 칸에만)
            tabs.append(tab); funcs.append((tab, func))
            r += 1
        end = r - 1
        blocks.append((start, end, tabs, funcs))
    for start, end, tabs, funcs in blocks:
        ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        merge_runs(ws, 2, start, end, lambda x: tabs[x - start])
        merge_runs(ws, 3, start, end, lambda x: funcs[x - start])


write(tc, True)
write(dev, False)
wb.calculation.fullCalcOnLoad = True
wb.save(dst)
print(f"테스트 케이스 {len(cases['common'])}건 + 엣지 {len(cases['edge'])}건 → {dst}")
